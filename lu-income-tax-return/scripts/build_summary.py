#!/usr/bin/env python3
"""Build the tax-return summary workbook from a JSON description.

Usage:
    python build_summary.py <data.json> <output.xlsx>

The JSON format is documented in assets/example_return.json. In short:

- "sections": list of {"title", "rows"}; each row has a unique "key", a "label",
  optional "p1" / "p2" (person 1 / person 2), "household", "form" (what MyGuichet
  shows), "source", "note", and "total": true for bold total rows.
- Values are numbers, or formulas: strings starting with "=" in which {key}
  refers to the same column of another row, and {key:p1} / {key:p2} / {key:hh} /
  {key:form} refer to a specific column. Example: "=MIN({decl},{cap})".
- If "household" is omitted and p1 or p2 is present, household = p1 + p2.
- "summary": rows for the Summary sheet, each {"label", "key", "per_person"}.
- "previous_year": {"year", "rows"}: rows with last year's p1/p2/household
  (numbers or {key} formulas within that sheet), "current": key of the matching
  Detail row, "why" and "source".
- "open_items": [{"priority", "topic", "found", "action"}].
- "attachments": [{"section", "file"}].
- "todo": list of strings shown on the Summary sheet.

Formulas are written without cached values: open and save the file in Excel or
LibreOffice (or run a recalculation script) before sharing it.
"""
import json
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

FONT = "Arial"
F_NORM = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_SEC = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_IN = Font(name=FONT, size=10, color="0000FF")      # typed from a document
F_LINK = Font(name=FONT, size=10, color="008000")    # link to another sheet
F_SMALL = Font(name=FONT, size=9, italic=True, color="595959")
FILL_SEC = PatternFill("solid", fgColor="1F4E78")
FILL_HEAD = PatternFill("solid", fgColor="D9E1F2")
FILL_TOT = PatternFill("solid", fgColor="F2F2F2")
FILL_WARN = PatternFill("solid", fgColor="FFFF00")
NUM = '#,##0.00;[Red]-#,##0.00;-'
PCT = '0.0%;[Red]-0.0%;-'
WRAP = Alignment(wrap_text=True, vertical="top")

COLS = {"p1": "B", "p2": "C", "hh": "D", "form": "E"}
PLACEHOLDER = re.compile(r"\{([A-Za-z0-9_]+)(?::(p1|p2|hh|form))?\}")


def put(ws, r, c, v, font=F_NORM, fmt=None, fill=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    return cell


def resolve(value, col, rows, sheet_prefix=""):
    """Replace {key} / {key:col} placeholders by cell references."""
    if not (isinstance(value, str) and value.startswith("=")):
        return value

    def sub(m):
        key, explicit = m.group(1), m.group(2)
        if key not in rows:
            raise KeyError(f"Unknown row key in formula: {key!r} ({value})")
        letter = COLS[explicit] if explicit else col
        return f"{sheet_prefix}{letter}{rows[key]}"

    return PLACEHOLDER.sub(sub, value)


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def build(data, out):
    persons = data.get("persons", ["Person 1", "Person 2"])
    wb = Workbook()

    # ------------------------------------------------------------ Detail
    d = wb.active
    d.title = "Detail"
    put(d, 1, 1, data.get("title", "Income tax return – detail"), F_TITLE)
    put(d, 2, 1, "Blue = typed from a document · black = formula · column E = what the MyGuichet form shows.", F_SMALL)

    # pass 1: layout (row numbers per key)
    rows, layout, r = {}, [], 3
    for sec in data.get("sections", []):
        r += 1
        layout.append(("section", r, sec["title"]))
        r += 1
        layout.append(("head", r, None))
        r += 1
        for row in sec.get("rows", []):
            if row.get("subheading"):
                layout.append(("sub", r, row["subheading"]))
                r += 1
                continue
            key = row["key"]
            if key in rows:
                raise KeyError(f"Duplicate row key: {key}")
            rows[key] = r
            layout.append(("row", r, row))
            r += 1

    # pass 2: write
    heads = ["Item", f"{persons[0]} (EUR)", f"{persons[1]} (EUR)" if len(persons) > 1 else "",
             "Household (EUR)", "Shown in MyGuichet (EUR)", "Source document", "Note"]
    for kind, rr, item in layout:
        if kind == "section":
            for c in range(1, 8):
                put(d, rr, c, item if c == 1 else None, F_SEC, fill=FILL_SEC)
        elif kind == "head":
            for c, h in enumerate(heads, 1):
                put(d, rr, c, h, F_BOLD, fill=FILL_HEAD, align=WRAP)
        elif kind == "sub":
            put(d, rr, 1, item, F_BOLD)
        else:
            row = item
            total = row.get("total", False)
            fill = FILL_TOT if total else None
            put(d, rr, 1, row["label"], F_BOLD if total else F_NORM, fill=fill, align=WRAP)
            for c, k in ((2, "p1"), (3, "p2")):
                v = row.get(k)
                if v is None:
                    put(d, rr, c, None, fill=fill)
                else:
                    font = F_BOLD if total else (F_NORM if is_formula(v) else F_IN)
                    put(d, rr, c, resolve(v, COLS[k], rows), font, NUM, fill)
            if "household" in row:
                hv = row["household"]
            elif row.get("p1") is not None or row.get("p2") is not None:
                hv = f"=SUM(B{rr}:C{rr})"
            else:
                hv = None
            if hv is None:
                put(d, rr, 4, None, fill=fill)
            else:
                font = F_BOLD if total else (F_NORM if is_formula(hv) else F_IN)
                put(d, rr, 4, resolve(hv, "D", rows), font, row.get("format", NUM), fill)
            fv = row.get("form")
            put(d, rr, 5, resolve(fv, "E", rows), F_IN if fv is not None else F_NORM, NUM, fill)
            put(d, rr, 6, row.get("source", ""), F_NORM, fill=fill, align=WRAP)
            put(d, rr, 7, row.get("note", ""), F_NORM, fill=FILL_WARN if row.get("warn") else fill, align=WRAP)
    for col, w in zip("ABCDEFG", (56, 15, 15, 17, 17, 48, 60)):
        d.column_dimensions[col].width = w
    d.freeze_panes = "B4"

    # ------------------------------------------------------------ Summary
    s = wb.create_sheet("Summary", 0)
    put(s, 1, 1, data.get("title", "Income tax return – summary"), F_TITLE)
    put(s, 2, 1, data.get("intro", ""), F_NORM)
    put(s, 3, 1, data.get("status", ""), F_BOLD)
    r = 5
    for c, h in enumerate(["Key figure (EUR)", persons[0], persons[1] if len(persons) > 1 else "",
                           "Household", "Shown in MyGuichet", "Difference"], 1):
        put(s, r, c, h, F_BOLD, fill=FILL_HEAD, align=WRAP)
    for item in data.get("summary", []):
        r += 1
        k = rows[item["key"]]
        put(s, r, 1, item["label"], F_BOLD if item.get("bold") else F_NORM)
        if item.get("per_person", True):
            put(s, r, 2, f"=Detail!B{k}", F_LINK, NUM)
            put(s, r, 3, f"=Detail!C{k}", F_LINK, NUM)
        put(s, r, 4, f"=Detail!D{k}", F_LINK, NUM)
        put(s, r, 5, f"=Detail!E{k}", F_LINK, NUM)
        put(s, r, 6, f'=IF(Detail!E{k}="","",ROUND(D{r}-E{r},2))', F_NORM, NUM)
    r += 2
    put(s, r, 1, "Colour code: blue = typed from a document, black = formula, green = link to the Detail sheet, "
                 "yellow = needs attention.", F_SMALL)
    if data.get("todo"):
        r += 2
        put(s, r, 1, "What is left for you", F_BOLD, fill=FILL_HEAD)
        for c in range(2, 7):
            put(s, r, c, None, fill=FILL_HEAD)
        for t in data["todo"]:
            r += 1
            put(s, r, 1, t)
    s.column_dimensions["A"].width = 62
    for col in "BCDEF":
        s.column_dimensions[col].width = 18

    # ------------------------------------------------------------ Previous year
    py = data.get("previous_year")
    if py:
        c = wb.create_sheet("Previous year", 1)
        year = py.get("year", "previous")
        cur = data.get("tax_year", "current")
        put(c, 1, 1, f"{year} (as filed) compared with {cur}", F_TITLE)
        put(c, 2, 1, f"{year} = typed from last year's filed return (blue). {cur} = links to the Detail sheet (green).", F_SMALL)
        hdr = ["Item (EUR)", f"{year} {persons[0]}", f"{year} {persons[1] if len(persons) > 1 else ''}",
               f"{year} household", f"{cur} {persons[0]}", f"{cur} {persons[1] if len(persons) > 1 else ''}",
               f"{cur} household", "Change", "Change %", "What explains it", f"{year} source"]
        for ci, h in enumerate(hdr, 1):
            put(c, 4, ci, h, F_BOLD, fill=FILL_HEAD, align=WRAP)
        prow, r = {}, 5
        for row in py.get("rows", []):
            if row.get("key"):
                prow[row["key"]] = r
            r += 1
        r = 5
        pcols = {"p1": "B", "p2": "C", "hh": "D", "form": "D"}
        for row in py.get("rows", []):
            total = row.get("total", False)
            fill = FILL_TOT if total else None
            put(c, r, 1, row["label"], F_BOLD if total else F_NORM, fill=fill, align=WRAP)

            def prev_val(v, col):
                if not is_formula(v):
                    return v

                def sub(m):
                    key, explicit = m.group(1), m.group(2)
                    return f"{pcols[explicit] if explicit else col}{prow[key]}"
                return PLACEHOLDER.sub(sub, v)

            for ci, k, col in ((2, "p1", "B"), (3, "p2", "C")):
                v = row.get(k)
                if v is not None:
                    put(c, r, ci, prev_val(v, col), F_BOLD if total else (F_NORM if is_formula(v) else F_IN), NUM, fill)
                else:
                    put(c, r, ci, None, fill=fill)
            hv = row.get("household", f"=SUM(B{r}:C{r})" if (row.get("p1") is not None or row.get("p2") is not None) else None)
            put(c, r, 4, prev_val(hv, "D"), F_BOLD if total else (F_NORM if is_formula(hv) else F_IN), NUM, fill)
            ck = row.get("current")
            if ck:
                k = rows[ck]
                if row.get("per_person", True):
                    put(c, r, 5, f"=Detail!B{k}", F_LINK, NUM, fill)
                    put(c, r, 6, f"=Detail!C{k}", F_LINK, NUM, fill)
                else:
                    put(c, r, 5, None, fill=fill)
                    put(c, r, 6, None, fill=fill)
                put(c, r, 7, f"=Detail!D{k}", F_LINK, NUM, fill)
            put(c, r, 8, f"=G{r}-D{r}", F_BOLD if total else F_NORM, NUM, fill)
            put(c, r, 9, f'=IF(N(D{r})=0,"",H{r}/ABS(D{r}))', F_BOLD if total else F_NORM, PCT, fill)
            put(c, r, 10, row.get("why", ""), F_NORM, fill=fill, align=WRAP)
            put(c, r, 11, row.get("source", ""), F_SMALL, fill=fill, align=WRAP)
            r += 1
        c.column_dimensions["A"].width = 44
        for col in "BCDEFGH":
            c.column_dimensions[col].width = 15
        c.column_dimensions["I"].width = 10
        c.column_dimensions["J"].width = 70
        c.column_dimensions["K"].width = 26
        c.freeze_panes = "B5"

    # ------------------------------------------------------------ Open items
    o = wb.create_sheet("Open items")
    put(o, 1, 1, "Open items and choices to confirm before signing", F_TITLE)
    for ci, h in enumerate(["#", "Priority", "Topic", "What was found / entered", "What you need to do"], 1):
        put(o, 3, ci, h, F_BOLD, fill=FILL_HEAD, align=WRAP)
    for i, it in enumerate(data.get("open_items", []), 1):
        rr = 3 + i
        put(o, rr, 1, i, F_NORM, align=WRAP)
        put(o, rr, 2, it.get("priority", ""), F_BOLD, fill=FILL_WARN if it.get("priority") == "High" else None, align=WRAP)
        put(o, rr, 3, it.get("topic", ""), F_BOLD, align=WRAP)
        put(o, rr, 4, it.get("found", ""), F_NORM, align=WRAP)
        put(o, rr, 5, it.get("action", ""), F_NORM, align=WRAP)
    for col, w in zip("ABCDE", (4, 9, 30, 85, 70)):
        o.column_dimensions[col].width = w
    o.freeze_panes = "A4"

    # ------------------------------------------------------------ Attachments
    a = wb.create_sheet("Attachments")
    put(a, 1, 1, "Supporting documents to upload with the return", F_TITLE)
    for ci, h in enumerate(["#", "Section", "File"], 1):
        put(a, 3, ci, h, F_BOLD, fill=FILL_HEAD)
    for i, it in enumerate(data.get("attachments", []), 1):
        put(a, 3 + i, 1, i)
        put(a, 3 + i, 2, it.get("section", ""))
        put(a, 3 + i, 3, it.get("file", ""))
    a.column_dimensions["A"].width = 5
    a.column_dimensions["B"].width = 36
    a.column_dimensions["C"].width = 90

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    with open(sys.argv[1], encoding="utf-8") as fh:
        data = json.load(fh)
    build(data, sys.argv[2])
    print(f"saved {sys.argv[2]}")


if __name__ == "__main__":
    main()
